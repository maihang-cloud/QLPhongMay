from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file, Response
from werkzeug.middleware.proxy_fix import ProxyFix
from database import get_db_connection, init_db
from report_exporter import (
    get_schedule_ca_slots,
    build_semester_report_excel,
    summarize_available_semesters,
    filter_schedules_for_report
)
from functools import wraps
import urllib.request
import urllib.parse
import secrets
import json
import os

app = Flask(__name__)

# Hỗ trợ nhận diện đúng giao thức HTTPS phía sau Reverse Proxy của PythonAnywhere
app.wsgi_app = ProxyFix(app.wsgi_app, x_proto=1, x_host=1)

# ==========================================
# CẤU HÌNH BẢO MẬT & GOOGLE OAUTH 2.0
# ==========================================
# Điền trực tiếp GOOGLE_CLIENT_ID và GOOGLE_CLIENT_SECRET tại đây hoặc cấu hình trong file WSGI trên PythonAnywhere
GOOGLE_CLIENT_ID = os.environ.get('GOOGLE_CLIENT_ID', '').strip()
GOOGLE_CLIENT_SECRET = os.environ.get('GOOGLE_CLIENT_SECRET', '').strip()

# Tự động đọc từ file "gg clent id code.txt" khi chạy thử ở máy cá nhân (file này đã được chặn trong .gitignore)
_local_secret_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'gg clent id code.txt')
if (not GOOGLE_CLIENT_ID or not GOOGLE_CLIENT_SECRET) and os.path.exists(_local_secret_file):
    try:
        with open(_local_secret_file, 'r', encoding='utf-8') as _f:
            for _line in _f:
                if _line.lower().startswith('client id:'):
                    GOOGLE_CLIENT_ID = _line.split(':', 1)[1].strip()
                elif _line.lower().startswith('client secret:'):
                    GOOGLE_CLIENT_SECRET = _line.split(':', 1)[1].strip()
    except Exception:
        pass

# Khai báo danh sách Quản trị viên
ADMIN_EMAILS = [
    "mlhang@hcmus.edu.vn",
    "tthan@hcmus.edu.vn",
    "nvtinh@hcmus.edu.vn"
]

# Tùy chọn: Giới hạn tên miền email được phép đăng nhập (Để danh sách rỗng [] nếu cho phép mọi tài khoản Google)
ALLOWED_EMAIL_DOMAINS = ["hcmus.edu.vn"]  # Ví dụ: ["hcmus.edu.vn", "student.hcmus.edu.vn"]


def get_current_user_email():
    """Lấy email người dùng đã xác thực từ Flask Session."""
    return session.get("user_email", "")


def get_current_user_name():
    """Lấy tên hiển thị của người dùng đã xác thực từ Flask Session."""
    email = get_current_user_email()
    name = session.get("user_name", "")
    if not name and email:
        return email.split("@")[0]
    return name


@app.context_processor
def inject_user_info():
    return dict(
        ADMIN_EMAILS=ADMIN_EMAILS,
        current_user_email=get_current_user_email(),
        current_user_name=get_current_user_name()
    )


def login_required(f):
    """Yêu cầu đăng nhập Google trước khi thực hiện thao tác."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not get_current_user_email():
            if request.is_json or request.path.startswith('/submit_bulk_booking'):
                return jsonify({'success': False, 'message': 'Vui lòng đăng nhập bằng tài khoản Google trước khi gửi yêu cầu.'}), 401
            flash("Vui lòng đăng nhập bằng tài khoản Google để tiếp tục.", "warning")
            return redirect(url_for('login_page', next=request.url))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    """Yêu cầu tài khoản Google thuộc danh sách ADMIN_EMAILS."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        user_email = get_current_user_email()
        if not user_email:
            flash("Vui lòng đăng nhập bằng tài khoản Google Quản trị viên.", "warning")
            return redirect(url_for('login_page', next=request.url))
        if user_email not in ADMIN_EMAILS:
            flash(f"Bạn không có quyền quản trị! (Tài khoản hiện tại: {user_email})", "danger")
            return redirect(url_for('home_page'))
        return f(*args, **kwargs)
    return decorated_function


# Thiết lập secret_key để mã hóa session
app.secret_key = os.environ.get('SECRET_KEY', 'KHTN_TKB_Secret_Key_@2026')
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024  # Giới hạn upload tối đa 10MB
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.jinja_env.auto_reload = True

# Khởi tạo DB khi app bắt đầu chạy
with app.app_context():
    init_db()


# ==========================================
# PHÂN HỆ XÁC THỰC GOOGLE OAUTH 2.0
# ==========================================

@app.route('/login')
def login_page():
    next_url = request.args.get('next') or url_for('home_page')
    if get_current_user_email():
        return redirect(next_url)
    oauth_configured = bool(GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET)
    return render_template('login.html', next_url=next_url, oauth_configured=oauth_configured)


@app.route('/login/google')
def google_login():
    next_url = request.args.get('next') or url_for('home_page')
    session['next_url'] = next_url

    if not GOOGLE_CLIENT_ID or not GOOGLE_CLIENT_SECRET:
        flash("Hệ thống chưa được cấu hình GOOGLE_CLIENT_ID và GOOGLE_CLIENT_SECRET.", "danger")
        return redirect(url_for('login_page', next=next_url))

    state = secrets.token_urlsafe(24)
    session['oauth_state'] = state
    redirect_uri = url_for('google_callback', _external=True)

    params = {
        'client_id': GOOGLE_CLIENT_ID,
        'redirect_uri': redirect_uri,
        'response_type': 'code',
        'scope': 'openid email profile',
        'state': state,
        'prompt': 'select_account'
    }
    auth_url = 'https://accounts.google.com/o/oauth2/v2/auth?' + urllib.parse.urlencode(params)
    return redirect(auth_url)


@app.route('/auth/google/callback')
def google_callback():
    error = request.args.get('error')
    if error:
        flash(f"Đăng nhập Google bị hủy hoặc gặp lỗi: {error}", "danger")
        return redirect(url_for('login_page'))

    state = request.args.get('state')
    if not state or state != session.pop('oauth_state', None):
        flash("Phiên xác thực không hợp lệ (CSRF State Mismatch). Vui lòng thử lại.", "danger")
        return redirect(url_for('login_page'))

    code = request.args.get('code')
    if not code:
        flash("Không nhận được mã xác thực từ Google.", "danger")
        return redirect(url_for('login_page'))

    redirect_uri = url_for('google_callback', _external=True)
    token_data = urllib.parse.urlencode({
        'code': code,
        'client_id': GOOGLE_CLIENT_ID,
        'client_secret': GOOGLE_CLIENT_SECRET,
        'redirect_uri': redirect_uri,
        'grant_type': 'authorization_code'
    }).encode('utf-8')

    try:
        token_req = urllib.request.Request(
            'https://oauth2.googleapis.com/token',
            data=token_data,
            headers={'Content-Type': 'application/x-www-form-urlencoded'}
        )
        with urllib.request.urlopen(token_req, timeout=10) as resp:
            token_json = json.loads(resp.read().decode('utf-8'))

        access_token = token_json.get('access_token')
        if not access_token:
            flash("Không lấy được Access Token từ Google.", "danger")
            return redirect(url_for('login_page'))

        userinfo_req = urllib.request.Request(
            'https://www.googleapis.com/oauth2/v3/userinfo',
            headers={'Authorization': f'Bearer {access_token}'}
        )
        with urllib.request.urlopen(userinfo_req, timeout=10) as resp:
            userinfo = json.loads(resp.read().decode('utf-8'))

        email = (userinfo.get('email') or '').strip().lower()
        name = (userinfo.get('name') or '').strip()

        if not email:
            flash("Tài khoản Google không cung cấp địa chỉ email hợp lệ.", "danger")
            return redirect(url_for('login_page'))

        if ALLOWED_EMAIL_DOMAINS:
            domain = email.split('@')[-1]
            if domain not in ALLOWED_EMAIL_DOMAINS and email not in ADMIN_EMAILS:
                flash(f"Vui lòng sử dụng email thuộc tên miền được phép ({', '.join(ALLOWED_EMAIL_DOMAINS)}).", "danger")
                return redirect(url_for('login_page'))

        session['user_email'] = email
        session['user_name'] = name or email.split('@')[0]

        next_url = session.pop('next_url', None) or url_for('home_page')
        flash(f"Xin chào {session['user_name']} ({email})!", "success")
        return redirect(next_url)

    except Exception as e:
        flash(f"Lỗi kết nối xác thực Google OAuth: {e}", "danger")
        return redirect(url_for('login_page'))


@app.route('/logout')
def logout():
    session.clear()
    flash("Đã đăng xuất khỏi hệ thống.", "success")
    return redirect(url_for('home_page'))


# ==========================================
# PHÂN HỆ USER (PUBLIC & AUTHENTICATED)
# ==========================================

import datetime
import re

ROOMS_LIST = ['5.1', '5.2', '5.3', '6.1', '6.2', '6.3', '6.5', '6.6']

def parse_study_period_date(raw_val):
    """Phân tích chuỗi Giai_Doan (VD: '28/09/2026', '2026-09-28 00:00:00') thành datetime.date."""
    if not raw_val:
        return None
    if isinstance(raw_val, datetime.datetime):
        return raw_val.date()
    if isinstance(raw_val, datetime.date):
        return raw_val
    s = str(raw_val).strip()
    if not s:
        return None
    # Thử định dạng YYYY-MM-DD (hoặc YYYY-MM-DD HH:MM:SS từ Excel)
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%Y/%m/%d'):
        try:
            part = s[:10] if fmt in ('%Y-%m-%d', '%Y/%m/%d') and len(s) >= 10 else s
            return datetime.datetime.strptime(part, fmt).date()
        except ValueError:
            pass
    # Tìm mẫu DD/MM/YYYY hoặc DD-MM-YYYY trong chuỗi bất kỳ
    m = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})', s)
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass
    return None

def normalize_study_period_str(raw_val):
    """Chuẩn hóa giá trị Giai_Doan về chuỗi DD/MM/YYYY nếu là ngày hợp lệ."""
    dt = parse_study_period_date(raw_val)
    if dt:
        return dt.strftime('%d/%m/%Y')
    return str(raw_val).strip() if raw_val else ''

def enrich_schedule_row(row, ref_date=None):
    """Bổ sung thông tin chu kỳ lặp (Tuần X/N, Ngày kết thúc, Trạng thái Đóng/Mở cửa) cho 1 dòng lịch học."""
    if ref_date is None:
        ref_date = datetime.date.today()
    d = dict(row)
    num_weeks = int(d.get('num_weeks') or 15)
    is_ended_early = int(d.get('is_ended_early') or 0)
    ended_early_date = d.get('ended_early_date') or ''
    raw_period = d.get('study_period') or ''

    start_dt = parse_study_period_date(raw_period)
    ref_monday = ref_date - datetime.timedelta(days=ref_date.weekday())

    d['num_weeks'] = num_weeks
    d['is_ended_early'] = is_ended_early
    d['ended_early_date'] = ended_early_date

    if start_dt:
        start_monday = start_dt - datetime.timedelta(days=start_dt.weekday())
        end_dt = start_monday + datetime.timedelta(days=num_weeks * 7 - 1)
        week_num = ((ref_monday - start_monday).days // 7) + 1

        d['start_date_str'] = start_dt.strftime('%d/%m/%Y')
        d['end_date_str'] = end_dt.strftime('%d/%m/%Y')
        d['week_num'] = week_num
        d['has_valid_date'] = True
        d['is_in_selected_week'] = (1 <= week_num <= num_weeks)

        # Tính tuần dừng nếu kết thúc sớm
        stopped_week = week_num
        if ended_early_date:
            early_dt = parse_study_period_date(ended_early_date)
            if early_dt:
                early_monday = early_dt - datetime.timedelta(days=early_dt.weekday())
                stopped_week = max(1, min(num_weeks, ((early_monday - start_monday).days // 7) + 1))
        d['stopped_week'] = max(1, min(num_weeks, stopped_week))

        if is_ended_early:
            d['door_status'] = 'ended_early'
            d['door_open'] = False
        elif week_num < 1:
            d['door_status'] = 'upcoming'
            d['door_open'] = False
        elif week_num > num_weeks:
            d['door_status'] = 'completed'
            d['door_open'] = False
        else:
            d['door_status'] = 'active'
            d['door_open'] = True
    else:
        d['start_date_str'] = raw_period if raw_period else 'Thường xuyên'
        d['end_date_str'] = ''
        d['week_num'] = 1
        d['stopped_week'] = 1
        d['has_valid_date'] = False
        d['is_in_selected_week'] = True
        if is_ended_early:
            d['door_status'] = 'ended_early'
            d['door_open'] = False
        else:
            d['door_status'] = 'active'
            d['door_open'] = True

    return d

def get_detailed_free_slots_data(day_of_week, ref_date=None):
    """Tính toán chi tiết các khoảng phòng trống dạng structured data cho ngày chỉ định."""
    conn = get_db_connection()
    schedules = conn.execute(
        'SELECT * FROM schedules WHERE day_of_week = ? AND COALESCE(is_ended_early, 0) = 0 ORDER BY start_period ASC',
        (day_of_week,)
    ).fetchall()
    conn.close()
    
    room_schedules = {room: [] for room in ROOMS_LIST}
    for row in schedules:
        enriched = enrich_schedule_row(row, ref_date=ref_date)
        # Nếu môn chưa đến tuần bắt đầu (upcoming), đã hết chu kỳ tuần lặp (completed), hoặc đã kết thúc sớm (ended_early)
        # thì giải phóng phòng trống cho ngày/tuần đang tra cứu
        if not enriched['door_open']:
            continue
        room = row['room']
        if room in room_schedules:
            room_schedules[room].append(row)
            
    def parse_slots(slots):
        res = []
        for s, e in slots:
            start_tiet = s + 1.0
            if e.is_integer():
                end_tiet = e
            else:
                end_tiet = e + 1.0
            st_str = f"{int(start_tiet)}" if start_tiet.is_integer() else f"{start_tiet}"
            ed_str = f"{int(end_tiet)}" if end_tiet.is_integer() else f"{end_tiet}"
            st_val = int(start_tiet) if start_tiet.is_integer() else start_tiet
            ed_val = int(end_tiet) if end_tiet.is_integer() else end_tiet
            res.append({
                'start_period': st_val,
                'end_period': ed_val,
                'label': f"Tiết {st_str} - {ed_str}"
            })
        return res

    morning = []
    afternoon = []
    
    for room, classes in room_schedules.items():
        # Sáng (0.0 đến 5.0)
        m_slots = []
        t_trong = 0.0
        for c in classes:
            start = float(c['start_period'])
            end = float(c['end_period'])
            if start > 5.0: continue
            c_start_time = start - 1.0
            c_end_time = end if end.is_integer() else float(int(end)) - 0.5
            if c_start_time > t_trong:
                if (c_start_time - t_trong) >= 2.5:
                    m_slots.append((t_trong, c_start_time))
            t_trong = max(t_trong, c_end_time)
        if 5.0 - t_trong >= 2.5:
            m_slots.append((t_trong, 5.0))
        for item in parse_slots(m_slots):
            morning.append({'room': room, **item})

        # Chiều (5.0 đến 10.0)
        a_slots = []
        t_trong = 5.0
        for c in classes:
            start = float(c['start_period'])
            end = float(c['end_period'])
            if end <= 5.0: continue
            c_start_time = max(start - 1.0, 5.0)
            c_end_time = end if end.is_integer() else float(int(end)) - 0.5
            if c_start_time > t_trong:
                if (c_start_time - t_trong) >= 2.5:
                    a_slots.append((t_trong, c_start_time))
            t_trong = max(t_trong, c_end_time)
        if 10.0 - t_trong >= 2.5:
            a_slots.append((t_trong, 10.0))
        for item in parse_slots(a_slots):
            afternoon.append({'room': room, **item})
            
    return morning, afternoon

def get_free_slots_for_day(day_of_week, ref_date=None):
    """Giữ nguyên tương thích ngược cho các nơi sử dụng định dạng dict, có hỗ trợ ref_date."""
    morning, afternoon = get_detailed_free_slots_data(day_of_week, ref_date=ref_date)
    morning_free = {}
    for item in morning:
        morning_free.setdefault(item['room'], []).append(item['label'])
    afternoon_free = {}
    for item in afternoon:
        afternoon_free.setdefault(item['room'], []).append(item['label'])
    return morning_free, afternoon_free

@app.route('/')
def index():
    """Chuyển hướng đến trang chủ"""
    return redirect(url_for('home_page'))

@app.route('/bao-cao-su-co')
@login_required
def report_page():
    return render_template('report.html')

@app.route('/submit_issue', methods=['POST'])
@login_required
def submit_issue():
    reporter_name = request.form.get('reporter_name') or get_current_user_name()
    department = request.form.get('department')
    room = request.form.get('room')
    machine_number = request.form.get('machine_number')
    issue_category = request.form.get('issue_category')
    description = request.form.get('description')
    
    conn = get_db_connection()
    try:
        user_email = get_current_user_email()
        conn.execute('''
        INSERT INTO issue_reports (reporter_name, department, room, machine_number, issue_category, description, email)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    ''', (reporter_name, department, room, machine_number, issue_category, description, user_email))
        conn.commit()
        flash("Gửi báo cáo sự cố thành công! Bộ phận QTTB sẽ nhanh chóng xử lý.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"Có lỗi xảy ra: {e}", "error")
    finally:
        conn.close()
        
    return redirect(url_for('report_page'))

@app.route('/dang-ky-muon-phong')
@login_required
def booking_page():
    """Hiển thị trang chứa Form đăng ký mượn phòng."""
    selected_day = request.args.get('filter_day', 2, type=int)
    raw_date = request.args.get('filter_date', '').strip()
    ref_date = parse_study_period_date(raw_date) if raw_date else datetime.date.today()
    
    morning_free, afternoon_free = get_free_slots_for_day(selected_day, ref_date=ref_date)
    return render_template('booking.html', 
                           selected_day=selected_day, 
                           morning_free=morning_free, 
                           afternoon_free=afternoon_free)

@app.route('/api/free-slots')
def api_free_slots():
    """API trả về danh sách chi tiết các ca phòng trống theo ngày và mốc tuần (cho sidebar & popup chọn nhanh)."""
    day = request.args.get('day', 2, type=int)
    if day < 2 or day > 7:
        day = 2
    raw_date = request.args.get('date', '').strip()
    ref_date = parse_study_period_date(raw_date) if raw_date else None
    if ref_date is None:
        ref_date = datetime.date.today()

    # Tính đúng ngày của Thứ (day) trong tuần chứa ref_date
    monday_of_ref = ref_date - datetime.timedelta(days=ref_date.weekday())
    target_date = monday_of_ref + datetime.timedelta(days=day - 2)
    saturday_of_ref = monday_of_ref + datetime.timedelta(days=5)

    morning, afternoon = get_detailed_free_slots_data(day, ref_date=target_date)
    return jsonify({
        'day': day,
        'target_date': target_date.strftime('%d/%m/%Y'),
        'target_date_iso': target_date.strftime('%Y-%m-%d'),
        'week_range': f"{monday_of_ref.strftime('%d/%m/%Y')} – {saturday_of_ref.strftime('%d/%m/%Y')}",
        'morning': morning,
        'afternoon': afternoon
    })

@app.route('/trang-chu')
def home_page():
    """Hiển thị trang chủ phòng máy."""
    return render_template('home.html')

@app.route('/booking', methods=['POST'])
@login_required
def submit_booking():
    """Xử lý data khi User gửi form mượn phòng."""
    requester_name = request.form.get('requester_name') or get_current_user_name()
    room = request.form.get('room')
    day_of_week = int(request.form.get('day_of_week')) if request.form.get('day_of_week') else None
    start_period = float(request.form.get('start_period')) if request.form.get('start_period') else None
    end_period = float(request.form.get('end_period')) if request.form.get('end_period') else None
    purpose = request.form.get('purpose')
    subject_code = request.form.get('subject_code')
    class_code = request.form.get('class_code')
    subject_name = request.form.get('subject_name')
    department = request.form.get('department')
    study_period = request.form.get('study_period') or request.form.get('study_date') or ''
    software_required = request.form.get('software_required')
    
    # Bug #9: Validate tiết bắt đầu và kết thúc
    if start_period is not None and end_period is not None:
        if start_period >= end_period:
            flash("Tiết bắt đầu phải nhỏ hơn tiết kết thúc!", "danger")
            return redirect(url_for('booking_page'))
        if start_period < 1 or end_period > 10:
            flash("Tiết phải nằm trong khoảng từ 1 đến 10!", "danger")
            return redirect(url_for('booking_page'))

    conn = get_db_connection()
    user_email = get_current_user_email()
    conn.execute('''
        INSERT INTO booking_requests (requester_name, room, day_of_week, start_period, end_period, purpose, subject_code, class_code, subject_name, department, study_period, software_required, status, email, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pending', ?, datetime('now', 'localtime'))
    ''', (requester_name, room, day_of_week, start_period, end_period, purpose, subject_code, class_code, subject_name, department, study_period, software_required, user_email))
    conn.commit()
    conn.close()
    
    flash("Gửi yêu cầu mượn phòng thành công! Vui lòng chờ Bộ phận Quản trị thiết bị duyệt.", "success")
    return redirect(url_for('booking_page'))

def is_pending_request_active_on_date(pending_row, ref_date):
    """Kiểm tra yêu cầu đang chờ duyệt (Pending) có trùng giai đoạn tuần với ngày đăng ký mới hay không."""
    if not ref_date:
        return True
    p_dt = parse_study_period_date(pending_row['study_period'] if 'study_period' in pending_row.keys() else '')
    if not p_dt:
        return True
    p_monday = p_dt - datetime.timedelta(days=p_dt.weekday())
    ref_monday = ref_date - datetime.timedelta(days=ref_date.weekday())
    week_diff = ((ref_monday - p_monday).days // 7) + 1
    return 1 <= week_diff <= 15

@app.route('/submit_bulk_booking', methods=['POST'])
@login_required
def submit_bulk_booking():
    """Xử lý data khi User gửi giỏ hàng đăng ký (nhiều môn)."""
    data = request.get_json()
    if not data or not isinstance(data, list):
        return jsonify({'success': False, 'message': 'Dữ liệu không hợp lệ'}), 400
    
    conn = get_db_connection()
    
    # 1. KIỂM TRA TRÙNG LỊCH VÀ TỰ ĐỘNG SẮP XẾP PHÒNG THEO ĐÚNG NGÀY/GIAI ĐOẠN CỦA TỪNG MÔN
    schedules_db = conn.execute(
        'SELECT room, day_of_week, start_period, end_period, study_period, num_weeks, is_ended_early, ended_early_date FROM schedules WHERE COALESCE(is_ended_early, 0) = 0'
    ).fetchall()
    pending_db = conn.execute(
        "SELECT room, day_of_week, start_period, end_period, study_period FROM booking_requests WHERE status = 'Pending'"
    ).fetchall()
    cart_assigned = []
    
    for item in data:
        room = item.get('room')
        try:
            dow = int(item.get('day_of_week'))
            req_start = float(item.get('start_period'))
            req_end = float(item.get('end_period'))
        except (ValueError, TypeError):
            conn.close()
            return jsonify({'success': False, 'message': 'Dữ liệu thời gian không hợp lệ'}), 400

        raw_item_date = item.get('study_period') or item.get('study_date') or ''
        item_ref_date = parse_study_period_date(raw_item_date) or datetime.date.today()

        active_for_item = [
            dict(s) for s in schedules_db
            if enrich_schedule_row(s, ref_date=item_ref_date)['door_open']
        ]
        pending_for_item = [
            dict(p) for p in pending_db
            if is_pending_request_active_on_date(p, item_ref_date)
        ]
        cart_for_item = [
            c for c in cart_assigned
            if is_pending_request_active_on_date(c, item_ref_date)
        ]
        schedules = active_for_item + pending_for_item + cart_for_item
            
        subject_name = item.get('subject_name', 'Môn học')
        rooms_to_check = [room] if room != 'Any' else ROOMS_LIST
        is_free_in_some_room = False
        assigned_room = None
        
        for r in rooms_to_check:
            overlap = False
            for s in schedules:
                if s['room'] == r and s['day_of_week'] == dow:
                    s_start = float(s['start_period'])
                    s_end = float(s['end_period'])
                    if req_start < s_end and req_end > s_start:
                        overlap = True
                        break
            if not overlap:
                is_free_in_some_room = True
                assigned_room = r
                break
                
        if not is_free_in_some_room:
            conn.close()
            error_msg = f"Môn '{subject_name}': Ca này đã được sử dụng (kể cả khi tìm phòng tự động), vui lòng chọn lịch khác."
            return jsonify({'success': False, 'message': error_msg}), 200
            
        # Gán phòng tự động nếu chọn 'Any'
        if room == 'Any':
            item['room'] = assigned_room
            
        # Thêm ca này vào danh sách cart_assigned để môn tiếp theo trong giỏ hàng không được xếp đè vào đây
        cart_assigned.append({
            'room': item['room'],
            'day_of_week': dow,
            'start_period': req_start,
            'end_period': req_end,
            'study_period': raw_item_date
        })

    # 2. TIẾN HÀNH LƯU VÀO DATABASE
    user_email = get_current_user_email()
    default_name = get_current_user_name()
    try:
        for item in data:
            conn.execute('''
                INSERT INTO booking_requests (
                    requester_name, room, day_of_week, start_period, end_period, purpose, 
                    subject_code, class_code, subject_name, department, study_period, software_required, 
                    status, email, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pending', ?, datetime('now', 'localtime'))
            ''', (
                item.get('requester_name') or default_name, 
                item.get('room'), 
                int(item.get('day_of_week')),
                float(item.get('start_period')), 
                float(item.get('end_period')), 
                item.get('purpose'), 
                item.get('subject_code'), 
                item.get('class_code'), 
                item.get('subject_name'), 
                item.get('department'), 
                item.get('study_period'), 
                item.get('software_required'),
                user_email
            ))
        conn.commit()
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500
    finally:
        conn.close()
    
    flash("Gửi danh sách yêu cầu mượn phòng thành công! Các phòng Tùy ý (nếu có) đã được hệ thống tự động xếp vào phòng trống.", "success")
    return jsonify({'success': True, 'redirect_url': url_for('booking_page')})

@app.route('/check_conflict', methods=['POST'])
def check_conflict():
    data = request.get_json()
    room = data.get('room')
    try:
        dow = int(data.get('day_of_week'))
        req_start = float(data.get('start_period'))
        req_end = float(data.get('end_period'))
    except (ValueError, TypeError):
        return jsonify({'is_free': False})

    raw_req_date = data.get('study_period') or data.get('study_date') or ''
    req_ref_date = parse_study_period_date(raw_req_date) or datetime.date.today()
        
    conn = get_db_connection()
    # Kiểm tra cả lịch đã duyệt (đang còn hiệu lực trong tuần của req_ref_date) và yêu cầu đang chờ duyệt
    schedules_confirmed = conn.execute(
        'SELECT room, day_of_week, start_period, end_period, study_period, num_weeks, is_ended_early, ended_early_date FROM schedules WHERE day_of_week = ? AND COALESCE(is_ended_early, 0) = 0', (dow,)
    ).fetchall()
    active_confirmed = [
        s for s in schedules_confirmed
        if enrich_schedule_row(s, ref_date=req_ref_date)['door_open']
    ]
    schedules_pending = conn.execute(
        'SELECT room, day_of_week, start_period, end_period, study_period FROM booking_requests WHERE day_of_week = ? AND status = ?',
        (dow, 'Pending')
    ).fetchall()
    active_pending = [
        p for p in schedules_pending
        if is_pending_request_active_on_date(p, req_ref_date)
    ]
    schedules = list(active_confirmed) + list(active_pending)
    conn.close()
    
    rooms_to_check = [room] if room != 'Any' else ROOMS_LIST
    
    is_free_in_some_room = False
    assigned_room = None
    for r in rooms_to_check:
        overlap = False
        for s in schedules:
            if s['room'] == r:
                s_start = float(s['start_period'])
                s_end = float(s['end_period'])
                if req_start < s_end and req_end > s_start:
                    overlap = True
                    break
        if not overlap:
            is_free_in_some_room = True
            assigned_room = r
            break
            
    return jsonify({'is_free': is_free_in_some_room, 'assigned_room': assigned_room})

@app.route('/thong-bao-nghi')
@login_required
def cancellation_page():
    """Hiển thị trang thông báo nghỉ gấp."""
    return render_template('cancellation.html')

@app.route('/submit-cancellation', methods=['POST'])
@login_required
def submit_cancellation():
    """Xử lý data khi giáo viên báo nghỉ gấp."""
    teacher_name = request.form.get('teacher_name') or get_current_user_name()
    room = request.form.get('room')
    cancel_date = request.form.get('cancel_date')
    period = request.form.get('period')
    reason = request.form.get('reason')
    
    conn = get_db_connection()
    user_email = get_current_user_email()
    conn.execute('''
        INSERT INTO cancellations (teacher_name, room, cancel_date, period, reason, email)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (teacher_name, room, cancel_date, period, reason, user_email))
    conn.commit()
    conn.close()
    
    flash("Gửi thông báo nghỉ thành công! Phòng QTTB sẽ tiến hành đóng cửa phòng để bảo vệ tài sản.", "success")
    return redirect(url_for('cancellation_page'))

# ==========================================
# PHÂN HỆ ADMIN (PROTECTED)
# ==========================================

@app.route('/admin/login')
@app.route('/admin')
def redirect_to_admin():
    if not get_current_user_email():
        return redirect(url_for('login_page', next=url_for('admin_panel')))
    return redirect(url_for('admin_panel'))

@app.route('/admin/panel')
@admin_required
def admin_panel():
    """Trang quản trị (Dashboard): Chỉ Admin mới vào được."""
    conn = get_db_connection()
    today = datetime.date.today()

    # 1. Lấy toàn bộ danh sách yêu cầu mượn phòng (ưu tiên Chờ xử lý lên đầu, lưu vết các yêu cầu đã xử lý)
    booking_requests = conn.execute("""
        SELECT * FROM booking_requests
        ORDER BY CASE status WHEN 'Pending' THEN 0 WHEN 'Processing' THEN 1 WHEN 'Approved' THEN 2 ELSE 3 END, id DESC
    """).fetchall()
    pending_requests = [r for r in booking_requests if r['status'] == 'Pending']

    # 2. Lấy danh sách lịch học hiện tại (để hiển thị trong bảng quản lý) và bổ sung thông tin chu kỳ tuần
    raw_schedules = conn.execute("SELECT * FROM schedules ORDER BY room ASC, day_of_week ASC, start_period ASC").fetchall()
    all_enriched_schedules = [enrich_schedule_row(r, ref_date=today) for r in raw_schedules]

    schedule_filter = request.args.get('schedule_filter', 'all')
    active_list = [s for s in all_enriched_schedules if s['door_open']]
    locked_list = [s for s in all_enriched_schedules if not s['door_open']]
    schedule_counts = {
        'all': len(all_enriched_schedules),
        'active': len(active_list),
        'locked': len(locked_list)
    }
    if schedule_filter == 'active':
        current_schedules = active_list
    elif schedule_filter == 'locked':
        current_schedules = locked_list
    else:
        current_schedules = all_enriched_schedules

    # Logic tìm phòng trống theo Thứ & Ngày/Tuần (hỗ trợ tra cứu Học kỳ 2)
    filter_day = request.args.get('filter_day', 2, type=int)
    raw_free_date = request.args.get('filter_date', '').strip()
    free_ref_date = parse_study_period_date(raw_free_date) if raw_free_date else today
    if not free_ref_date:
        free_ref_date = today
    free_monday = free_ref_date - datetime.timedelta(days=free_ref_date.weekday())
    free_target_date = free_monday + datetime.timedelta(days=filter_day - 2)
    filter_date_str = free_target_date.strftime('%Y-%m-%d')
    free_target_date_display = free_target_date.strftime('%d/%m/%Y')
    admin_morning_free, admin_afternoon_free = get_free_slots_for_day(filter_day, ref_date=free_target_date)

    # 3. Lấy dữ liệu hiển thị dạng Lưới theo tuần (Chế độ Google Calendar)
    rooms = ['5.1', '5.2', '5.3', '6.1', '6.2', '6.3', '6.5', '6.6']
    selected_room = request.args.get('room', '5.1')

    raw_tt_date = request.args.get('timetable_date', '').strip()
    selected_date = parse_study_period_date(raw_tt_date) if raw_tt_date else today
    if not selected_date:
        selected_date = today

    monday_date = selected_date - datetime.timedelta(days=selected_date.weekday())
    saturday_date = monday_date + datetime.timedelta(days=5)
    prev_week_date_str = (monday_date - datetime.timedelta(days=7)).strftime('%Y-%m-%d')
    next_week_date_str = (monday_date + datetime.timedelta(days=7)).strftime('%Y-%m-%d')
    timetable_date_str = selected_date.strftime('%Y-%m-%d')
    timetable_date_display = selected_date.strftime('%d/%m/%Y')
    today_date_str = today.strftime('%Y-%m-%d')
    week_range_label = f"Thứ Hai {monday_date.strftime('%d/%m')} – Thứ Bảy {saturday_date.strftime('%d/%m/%Y')}"

    day_headers = {}
    for d in range(2, 8):
        d_date = monday_date + datetime.timedelta(days=d - 2)
        day_headers[d] = {
            'date_str': d_date.strftime('%d/%m/%Y'),
            'short_date': d_date.strftime('%d/%m'),
            'is_today': (d_date == today)
        }

    schedules_grid = conn.execute('SELECT * FROM schedules WHERE room = ? ORDER BY start_period ASC', (selected_room,)).fetchall()
    week_data = {day: [] for day in range(2, 8)}

    for row in schedules_grid:
        day = row['day_of_week']
        if day in week_data:
            enriched = enrich_schedule_row(row, ref_date=monday_date)
            # Chỉ hiển thị trên Thời khóa biểu tuần này nếu học phần nằm trong chu kỳ lặp 1..num_weeks của tuần đang chọn
            if not enriched['is_in_selected_week']:
                continue

            start = float(enriched['start_period'])
            end = float(enriched['end_period'])

            start_time = start - 1.0
            if end.is_integer():
                end_time = end
            else:
                end_time = float(int(end)) - 0.5

            top_pct = start_time / 10 * 100
            height_pct = (end_time - start_time) / 10 * 100

            week_data[day].append({
                'data': enriched,
                'top': top_pct,
                'height': height_pct
            })

    cancellations = conn.execute("SELECT * FROM cancellations ORDER BY created_at DESC").fetchall()
    issues = conn.execute("SELECT * FROM issue_reports ORDER BY created_at DESC").fetchall()

    # 4. Tính toán số liệu thống kê cho Tab 📊 Dashboard & B42Lab NetAgent
    conn.execute("CREATE TABLE IF NOT EXISTS system_settings (key TEXT PRIMARY KEY, value TEXT)")
    conn.execute('''
        CREATE TABLE IF NOT EXISTS room_network_status (
            room TEXT PRIMARY KEY,
            target_ip TEXT DEFAULT '8.8.8.8',
            status TEXT DEFAULT 'online',
            ping_ms INTEGER DEFAULT 12,
            updated_at TEXT
        )
    ''')
    # Khởi tạo mã khóa bảo mật cho NetAgent nếu chưa có
    ak_row = conn.execute("SELECT value FROM system_settings WHERE key = 'netagent_key'").fetchone()
    if not ak_row or not ak_row['value']:
        default_agent_key = "b42lab-secret-2026"
        conn.execute("INSERT OR REPLACE INTO system_settings (key, value) VALUES ('netagent_key', ?)", (default_agent_key,))
        conn.commit()
        netagent_key = default_agent_key
    else:
        netagent_key = ak_row['value']

    # Đảm bảo đủ 8 phòng trong bảng room_network_status với đúng IP Router mặc định của trường
    default_router_ips = {
        '5.1': '10.29.37.53',
        '5.2': '10.29.37.54',
        '5.3': '10.29.37.52',
        '6.1': '10.29.37.60',
        '6.2': '10.29.37.62',
        '6.3': '10.29.37.64',
        '6.5': '10.29.37.63',
        '6.6': '10.29.37.61',
    }
    now_str = datetime.datetime.now().strftime('%H:%M:%S %d/%m/%Y')
    for idx, r in enumerate(rooms):
        existing_net = conn.execute("SELECT room, target_ip FROM room_network_status WHERE room = ?", (r,)).fetchone()
        real_ip = default_router_ips.get(r, '8.8.8.8')
        if not existing_net:
            conn.execute(
                "INSERT INTO room_network_status (room, target_ip, status, ping_ms, updated_at) VALUES (?, ?, 'online', ?, ?)",
                (r, real_ip, 46, now_str)
            )
        elif existing_net['target_ip'] and existing_net['target_ip'].startswith('192.168.1.'):
            # Tự động nâng cấp các IP mẫu 192.168.1.10x sang IP Router thực tế 10.29.37.xx
            conn.execute(
                "UPDATE room_network_status SET target_ip = ? WHERE room = ?",
                (real_ip, r)
            )
    conn.commit()

    net_rows = conn.execute("SELECT * FROM room_network_status").fetchall()
    net_map = {row['room']: dict(row) for row in net_rows}

    # Đếm số ca học theo từng phòng (theo chuẩn 1 ca = 2.5 tiết, tối đa 24 ca/tuần/phòng, 192 ca/tuần toàn hệ thống)
    room_stats = {}
    max_room_sessions = 24
    total_ca_used = 0
    day_ca_used = {d: 0 for d in range(2, 8)}

    for r in rooms:
        r_active = [s for s in all_enriched_schedules if s['room'] == r and s['door_open']]
        r_locked = [s for s in all_enriched_schedules if s['room'] == r and not s['door_open']]
        r_active_ca = sum(len(get_schedule_ca_slots(s['start_period'], s['end_period'])) for s in r_active)
        r_locked_ca = sum(len(get_schedule_ca_slots(s['start_period'], s['end_period'])) for s in r_locked)
        r_total_ca = r_active_ca + r_locked_ca
        if r_total_ca > max_room_sessions:
            max_room_sessions = r_total_ca
        for s in (r_active + r_locked):
            num_ca = len(get_schedule_ca_slots(s['start_period'], s['end_period']))
            total_ca_used += num_ca
            dow = int(s['day_of_week'])
            if dow in day_ca_used:
                day_ca_used[dow] += num_ca
        room_stats[r] = {
            'active': r_active_ca,
            'locked': r_locked_ca,
            'total': r_total_ca
        }

    for r in rooms:
        room_stats[r]['height_pct'] = int(round((room_stats[r]['total'] / max_room_sessions) * 85)) if room_stats[r]['total'] > 0 else 8

    # Công suất lấp đầy chuẩn (8 phòng x 24 ca/tuần = 192 ca/tuần; 1 ngày = 32 ca/toàn trường)
    utilization_pct = round((total_ca_used / 192.0) * 100, 2) if total_ca_used > 0 else 0
    peak_dow = max(day_ca_used, key=day_ca_used.get) if total_ca_used > 0 else 2
    peak_dow_pct = min(100, int(round((day_ca_used[peak_dow] / 32.0) * 100))) if total_ca_used > 0 else 0

    # Phân bổ theo Khoa / Đơn vị (Tự động nhận diện động & tính theo Số ca chuẩn 2.5 tiết)
    dept_groups = {}
    for s in all_enriched_schedules:
        raw_dept = (s.get('department') or '').strip()
        if not raw_dept:
            clean_dept = 'Khác'
        else:
            # Bỏ tiền tố "Khoa " ở đầu (không phân biệt hoa thường) để đồng nhất tên đơn vị
            clean_dept = re.sub(r'^[Kk]hoa\s+', '', raw_dept).strip()
            if not clean_dept:
                clean_dept = raw_dept
        norm_key = re.sub(r'[\s\-–—/]+', ' ', clean_dept.lower()).strip()
        num_ca = len(get_schedule_ca_slots(s['start_period'], s['end_period']))
        if norm_key not in dept_groups:
            dept_groups[norm_key] = {'labels': {}, 'total': 0}
        dept_groups[norm_key]['total'] += num_ca
        dept_groups[norm_key]['labels'][clean_dept] = dept_groups[norm_key]['labels'].get(clean_dept, 0) + num_ca

    sorted_depts = []
    for norm_key, info in dept_groups.items():
        best_label = max(info['labels'].items(), key=lambda x: (x[1], '-' in x[0], len(x[0])))[0]
        sorted_depts.append((best_label, info['total']))
    sorted_depts.sort(key=lambda x: (-x[1], x[0]))

    total_sched_ca = max(1, total_ca_used)
    palette = [
        '#2563eb',  # Xanh dương
        '#10b981',  # Xanh lá
        '#f59e0b',  # Cam vàng
        '#8b5cf6',  # Tím
        '#06b6d4',  # Xanh ngọc
        '#ec4899',  # Hồng
        '#f97316',  # Cam đậm
        '#14b8a6',  # Teal
        '#ef4444',  # Đỏ
        '#64748b',  # Xám xanh
    ]

    dept_items = []
    dept_counts = {}
    dept_pcts = {}
    gradient_stops = []
    current_deg = 0.0

    for idx, (label, count) in enumerate(sorted_depts):
        color = palette[idx % len(palette)]
        pct = max(1, int(round((count / total_sched_ca) * 100))) if count > 0 else 0
        deg_span = (count / total_sched_ca) * 360.0
        next_deg = 360.0 if idx == len(sorted_depts) - 1 else (current_deg + deg_span)
        gradient_stops.append(f"{color} {current_deg:.1f}deg {next_deg:.1f}deg")
        current_deg = next_deg

        dept_counts[label] = count
        dept_pcts[label] = pct
        dept_items.append({
            'name': label,
            'count': count,
            'pct': pct,
            'color': color
        })

    donut_gradient = f"conic-gradient({', '.join(gradient_stops)})" if gradient_stops else "conic-gradient(#e2e8f0 0deg 360deg)"

    # Trạng thái trực tiếp 8 phòng hôm nay & Lịch Mở/Khóa cửa hôm nay
    today_dow = today.weekday() + 2  # Thứ 2 = 2 .. Chủ nhật = 8
    display_dow = today_dow if 2 <= today_dow <= 7 else 2
    today_str_dmy = today.strftime('%d/%m/%Y')
    today_str_iso = today.strftime('%Y-%m-%d')

    # Kiểm tra các phòng có báo nghỉ hôm nay
    cancelled_rooms_today = set()
    for c in cancellations:
        c_date = (c['cancel_date'] or '').strip()
        if c_date in (today_str_dmy, today_str_iso):
            cancelled_rooms_today.add(str(c['room']).replace('Phòng ', '').replace('P.', '').strip())

    pending_issues_by_room = {r: 0 for r in rooms}
    for iss in issues:
        if iss['status'] != 'Resolved':
            r_code = str(iss['room']).strip()
            if r_code in pending_issues_by_room:
                pending_issues_by_room[r_code] += 1

    room_live_cards = []
    morning_open, morning_locked, morning_free_rooms = [], [], []
    afternoon_open, afternoon_locked, afternoon_free_rooms = [], [], []
    online_rooms_count = 0
    offline_rooms_list = []
    last_net_update = ''

    for r in rooms:
        r_today_classes = [s for s in all_enriched_schedules if s['room'] == r and int(s['day_of_week']) == display_dow and s['is_in_selected_week']]
        m_classes = [s for s in r_today_classes if float(s['start_period']) <= 5]
        a_classes = [s for s in r_today_classes if float(s['end_period']) >= 6]

        # Ca sáng
        if any(s['door_open'] for s in m_classes) and r not in cancelled_rooms_today:
            morning_open.append(f"P.{r}")
        elif m_classes or r in cancelled_rooms_today:
            morning_locked.append(f"P.{r}")
        else:
            morning_free_rooms.append(f"P.{r}")

        # Ca chiều
        if any(s['door_open'] for s in a_classes) and r not in cancelled_rooms_today:
            afternoon_open.append(f"P.{r}")
        elif a_classes or r in cancelled_rooms_today:
            afternoon_locked.append(f"P.{r}")
        else:
            afternoon_free_rooms.append(f"P.{r}")

        # Nhãn hoạt động trên thẻ Live Status 8 phòng (luôn ưu tiên hiển thị lịch học/trạng thái cửa)
        open_today = [s for s in r_today_classes if s['door_open']]
        ended_today = [s for s in r_today_classes if not s['door_open']]
        if r in cancelled_rooms_today:
            activity_label = "🔒 Khóa cửa (Lớp báo nghỉ)"
            activity_class = "text-danger fw-semibold"
        elif open_today:
            first_c = open_today[0]
            subj = first_c['subject_name']
            if len(subj) > 20:
                subj = subj[:19] + "…"
            activity_label = f"🔓 T.{first_c['start_period']:g}-{first_c['end_period']:g}: {subj}"
            activity_class = "text-primary fw-semibold"
        elif ended_today:
            activity_label = "🔒 Khóa cửa (Đã KT sớm)"
            activity_class = "text-secondary"
        else:
            activity_label = f"⚪ Trống lịch Thứ {display_dow}"
            activity_class = "text-muted"

        net_info = net_map.get(r, {'target_ip': '8.8.8.8', 'status': 'online', 'ping_ms': 12, 'updated_at': ''})
        net_status = net_info.get('status') or 'online'
        net_ping = net_info.get('ping_ms')
        if net_info.get('updated_at'):
            last_net_update = net_info['updated_at']

        if net_status == 'online':
            online_rooms_count += 1
        else:
            offline_rooms_list.append(f"P.{r}")

        room_live_cards.append({
            'room': r,
            'activity_label': activity_label,
            'activity_class': activity_class,
            'issue_count': pending_issues_by_room.get(r, 0),
            'target_ip': net_info.get('target_ip') or '8.8.8.8',
            'net_status': net_status,
            'ping_ms': net_ping,
            'updated_at': net_info.get('updated_at') or ''
        })

    dashboard_stats = {
        'room_stats': room_stats,
        'total_ca_used': total_ca_used,
        'total_ca_free': max(0, 192 - total_ca_used),
        'utilization_pct': utilization_pct,
        'peak_dow': peak_dow,
        'peak_dow_pct': peak_dow_pct,
        'dept_counts': dept_counts,
        'dept_pcts': dept_pcts,
        'dept_items': dept_items,
        'donut_gradient': donut_gradient,
        'room_live_cards': room_live_cards,
        'display_dow': display_dow,
        'morning_open': morning_open,
        'morning_locked': morning_locked,
        'morning_free_rooms': morning_free_rooms,
        'afternoon_open': afternoon_open,
        'afternoon_locked': afternoon_locked,
        'afternoon_free_rooms': afternoon_free_rooms,
        'online_rooms_count': online_rooms_count,
        'offline_rooms_list': offline_rooms_list,
        'last_net_update': last_net_update,
        'netagent_key': netagent_key
    }

    semester_summary = summarize_available_semesters(raw_schedules)

    # 5. Dữ liệu Quản lý Lịch trình & Chấm công Tăng ca (OT)
    conn.execute('''
        CREATE TABLE IF NOT EXISTS overtime_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ot_date TEXT NOT NULL,
            staff_name TEXT NOT NULL,
            role TEXT NOT NULL,
            start_time TEXT,
            end_time TEXT,
            hours REAL NOT NULL,
            purpose TEXT,
            status TEXT DEFAULT 'Pending',
            notes TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    ot_month = request.args.get('ot_month', 'all').strip()
    ot_filter_status = request.args.get('ot_status', 'all').strip()

    month_query = conn.execute("SELECT DISTINCT substr(ot_date, 1, 7) as m FROM overtime_records WHERE ot_date IS NOT NULL AND length(ot_date) >= 7 ORDER BY m DESC").fetchall()
    available_ot_months = [row['m'] for row in month_query if row['m']]
    current_month_str = today.strftime('%Y-%m')
    if current_month_str not in available_ot_months:
        available_ot_months.append(current_month_str)
        available_ot_months.sort(reverse=True)

    if not ot_month:
        ot_month = 'all'

    if ot_month == 'all':
        all_ot_rows = conn.execute("SELECT * FROM overtime_records ORDER BY ot_date DESC, id DESC").fetchall()
    else:
        all_ot_rows = conn.execute("SELECT * FROM overtime_records WHERE ot_date LIKE ? ORDER BY ot_date DESC, id DESC", (f"{ot_month}%",)).fetchall()

    ot_total_hours = sum(float(r['hours'] or 0) for r in all_ot_rows)
    ot_completed_hours = sum(float(r['hours'] or 0) for r in all_ot_rows if r['status'] == 'Completed')
    ot_completed_count = sum(1 for r in all_ot_rows if r['status'] == 'Completed')
    ot_pending_count = sum(1 for r in all_ot_rows if r['status'] == 'Pending')
    ot_staff_set = set(r['staff_name'] for r in all_ot_rows if r['staff_name'])

    if ot_filter_status == 'Completed':
        display_ot_rows = [r for r in all_ot_rows if r['status'] == 'Completed']
    elif ot_filter_status == 'Pending':
        display_ot_rows = [r for r in all_ot_rows if r['status'] == 'Pending']
    else:
        display_ot_rows = all_ot_rows

    ot_data = {
        'records': display_ot_rows,
        'selected_month': ot_month,
        'selected_status': ot_filter_status,
        'available_months': available_ot_months,
        'total_hours': round(ot_total_hours, 1),
        'completed_hours': round(ot_completed_hours, 1),
        'completed_count': ot_completed_count,
        'pending_count': ot_pending_count,
        'staff_count': len(ot_staff_set)
    }

    conn.close()
    return render_template(
        'admin_panel.html',
        booking_requests=booking_requests,
        pending_requests=pending_requests,
        schedules=current_schedules,
        schedule_counts=schedule_counts,
        schedule_filter=schedule_filter,
        rooms=rooms,
        selected_room=selected_room,
        week_data=week_data,
        day_headers=day_headers,
        timetable_date_str=timetable_date_str,
        timetable_date_display=timetable_date_display,
        today_date_str=today_date_str,
        prev_week_date_str=prev_week_date_str,
        next_week_date_str=next_week_date_str,
        week_range_label=week_range_label,
        cancellations=cancellations,
        filter_day=filter_day,
        filter_date_str=filter_date_str,
        free_target_date_display=free_target_date_display,
        morning_free=admin_morning_free,
        afternoon_free=admin_afternoon_free,
        issues=issues,
        dashboard_stats=dashboard_stats,
        semester_summary=semester_summary,
        ot_data=ot_data
    )

def measure_host_ping_ms(target_ip):
    """Đo ping (ms) tới 1 địa chỉ IP/Host trên cả Windows và Linux."""
    import subprocess
    import platform
    import time

    target_ip = (target_ip or '').strip()
    if not target_ip:
        return False, None

    is_win = platform.system().lower().startswith('win')
    cmd = ['ping', '-n', '1', '-w', '1000', target_ip] if is_win else ['ping', '-c', '1', '-W', '1', target_ip]
    try:
        t0 = time.time()
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=2.5)
        elapsed_ms = max(1, int(round((time.time() - t0) * 1000)))
        if res.returncode == 0:
            m = re.search(r'(?:time|thời gian)[=<](\d+)\s*ms', res.stdout, re.IGNORECASE)
            ping_val = int(m.group(1)) if m else elapsed_ms
            return True, ping_val
        return False, None
    except Exception:
        return False, None

@app.route('/api/network-heartbeat', methods=['POST'])
def api_network_heartbeat():
    """API nhận tín hiệu Heartbeat Push từ B42Lab NetAgent chạy tại phòng máy."""
    data = request.get_json(silent=True) or {}
    provided_key = (data.get('agent_key') or request.headers.get('X-Agent-Key') or '').strip()

    conn = get_db_connection()
    conn.execute("CREATE TABLE IF NOT EXISTS system_settings (key TEXT PRIMARY KEY, value TEXT)")
    ak_row = conn.execute("SELECT value FROM system_settings WHERE key = 'netagent_key'").fetchone()
    expected_key = ak_row['value'] if ak_row and ak_row['value'] else 'b42lab-secret-2026'

    if provided_key != expected_key:
        conn.close()
        return jsonify({'success': False, 'message': 'Invalid agent_key'}), 403

    now_str = datetime.datetime.now().strftime('%H:%M:%S %d/%m/%Y')
    results = data.get('rooms', {})
    # Hỗ trợ cả trường hợp gửi lẻ 1 phòng: {"room": "5.1", "status": "online", "ping_ms": 11}
    if not results and data.get('room'):
        results = {str(data['room']): {'status': data.get('status', 'online'), 'ping_ms': data.get('ping_ms')}}

    for r_code, r_info in results.items():
        if r_code in ['5.1', '5.2', '5.3', '6.1', '6.2', '6.3', '6.5', '6.6']:
            st = 'online' if r_info.get('status') in ('online', 1, True, '1') else 'offline'
            p_ms = r_info.get('ping_ms')
            conn.execute('''
                UPDATE room_network_status
                SET status = ?, ping_ms = ?, updated_at = ?
                WHERE room = ?
            ''', (st, p_ms, now_str, r_code))

    conn.commit()
    rows = conn.execute("SELECT room, target_ip FROM room_network_status").fetchall()
    targets = {row['room']: row['target_ip'] for row in rows}
    conn.close()

    return jsonify({
        'success': True,
        'updated_at': now_str,
        'targets': targets
    })

@app.route('/api/network-status')
@admin_required
def api_network_status():
    """API trả về trạng thái mạng 8 phòng máy theo thời gian thực cho Dashboard."""
    conn = get_db_connection()
    rows = conn.execute("SELECT room, target_ip, status, ping_ms, updated_at FROM room_network_status").fetchall()
    conn.close()

    room_map = {}
    online_count = 0
    offline_rooms = []
    last_update = ''

    for row in rows:
        r = row['room']
        st = row['status'] or 'online'
        if st == 'online':
            online_count += 1
        else:
            offline_rooms.append(f"P.{r}")
        if row['updated_at']:
            last_update = row['updated_at']
        room_map[r] = {
            'room': r,
            'target_ip': row['target_ip'],
            'status': st,
            'ping_ms': row['ping_ms'],
            'updated_at': row['updated_at'] or ''
        }

    return jsonify({
        'success': True,
        'online_count': online_count,
        'total_rooms': len(rows),
        'offline_rooms': offline_rooms,
        'last_update': last_update,
        'room_map': room_map
    })

@app.route('/admin/network-config', methods=['POST'])
@admin_required
def save_network_config():
    """Lưu cấu hình IP của 8 phòng máy từ giao diện Admin."""
    conn = get_db_connection()
    for r in ['5.1', '5.2', '5.3', '6.1', '6.2', '6.3', '6.5', '6.6']:
        ip_val = (request.form.get(f"ip_{r}") or '8.8.8.8').strip()
        conn.execute("UPDATE room_network_status SET target_ip = ? WHERE room = ?", (ip_val, r))
    conn.commit()
    conn.close()
    flash("✅ Đã lưu danh sách địa chỉ IP Router cho 8 phòng máy!", "success")
    return redirect(url_for('admin_panel', tab='dashboard'))

@app.route('/admin/network-ping-now', methods=['POST'])
@admin_required
def admin_network_ping_now():
    """Thực hiện Ping kiểm tra ngay lập tức tới 8 phòng máy."""
    conn = get_db_connection()
    rows = conn.execute("SELECT room, target_ip FROM room_network_status").fetchall()
    now_str = datetime.datetime.now().strftime('%H:%M:%S %d/%m/%Y')

    for row in rows:
        r = row['room']
        ip_addr = row['target_ip'] or '8.8.8.8'
        is_up, ping_ms = measure_host_ping_ms(ip_addr)
        st = 'online' if is_up else 'offline'
        conn.execute(
            "UPDATE room_network_status SET status = ?, ping_ms = ?, updated_at = ? WHERE room = ?",
            (st, ping_ms, now_str, r)
        )
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'updated_at': now_str})

@app.route('/admin/download-netagent')
@admin_required
def download_netagent_script():
    """Tải trực tiếp file script b42_network_agent.py đã điền sẵn 8 IP Router, URL và khóa bảo mật (trên RAM để không kích hoạt Flask reload)."""
    conn = get_db_connection()
    rows = conn.execute("SELECT room, target_ip FROM room_network_status").fetchall()
    conn.close()
    ip_dict = {row['room']: row['target_ip'] for row in rows}

    app_dir = os.path.dirname(os.path.abspath(__file__))
    agent_path = os.path.join(app_dir, 'b42_network_agent.py')
    if os.path.exists(agent_path):
        with open(agent_path, 'r', encoding='utf-8') as f:
            content = f.read()
        lines = ["ROOM_TARGETS = {"]
        for r in ['5.1', '5.2', '5.3', '6.1', '6.2', '6.3', '6.5', '6.6']:
            lines.append(f'    "{r}": "{ip_dict.get(r, "8.8.8.8")}",')
        lines.append("}")
        new_block = "\n".join(lines)
        content = re.sub(r'ROOM_TARGETS\s*=\s*\{[^}]*\}', new_block, content)
        current_server_url = request.host_url.rstrip('/')
        content = content.replace('DEFAULT_SERVER_URL = "http://127.0.0.1:5000"', f'DEFAULT_SERVER_URL = "{current_server_url}"')
        return Response(
            content,
            mimetype='application/octet-stream',
            headers={'Content-Disposition': 'attachment; filename="b42_network_agent.py"'}
        )
    return "File b42_network_agent.py not found", 404


@app.route('/admin/export-semester-report')
@admin_required
def export_semester_report():
    """Xuất file Excel (.xlsx) Báo cáo Tần suất Khai thác Phòng máy tính Học kỳ (4 Sheet chuẩn QTTB)."""
    semester_mode = request.args.get('semester', 'hk1').strip().lower()
    academic_year = request.args.get('academic_year', '').strip()
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()

    conn = get_db_connection()
    schedules_rows = conn.execute("SELECT * FROM schedules ORDER BY room ASC, day_of_week ASC, start_period ASC").fetchall()
    issues_rows = conn.execute("SELECT * FROM issue_reports ORDER BY created_at DESC").fetchall()
    bookings_rows = conn.execute("SELECT * FROM booking_requests ORDER BY id DESC").fetchall()
    try:
        net_rows = conn.execute("SELECT * FROM room_network_status").fetchall()
    except Exception:
        net_rows = []
    conn.close()

    filter_res = filter_schedules_for_report(
        all_schedules=schedules_rows,
        semester_mode=semester_mode,
        academic_year=academic_year,
        from_date_str=from_date,
        to_date_str=to_date
    )

    if len(filter_res['rows']) == 0:
        flash(
            f"⚠️ Không tìm thấy dữ liệu môn học nào thuộc {filter_res['semester_short']} ({filter_res['academic_year']}) "
            f"trong Thời khóa biểu hiện tại! Vui lòng kiểm tra lại thời gian Bắt đầu – Kết thúc (Giai đoạn) của các môn học "
            f"hoặc Import lịch học cho kỳ này trước khi xuất báo cáo.",
            "warning"
        )
        return redirect(url_for('admin_panel', tab='schedules'))

    excel_buffer, filename = build_semester_report_excel(
        all_schedules=schedules_rows,
        issues_list=issues_rows,
        bookings_list=bookings_rows,
        net_status_list=net_rows,
        semester_mode=semester_mode,
        academic_year=academic_year,
        from_date_str=from_date,
        to_date_str=to_date
    )

    return send_file(
        excel_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


def build_overtime_report_excel(ot_rows, month_str):
    """Xuất file Excel bảng chấm công và theo dõi tăng ca (OT)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Bang Cham Cong OT"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Calibri", size=14, bold=True, color="0F2942")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill("solid", fgColor="0F2942")
    fill_summary = PatternFill("solid", fgColor="F0FDFA")
    font_summary = Font(name="Calibri", size=11, bold=True, color="0D9488")
    font_body = Font(name="Calibri", size=10, color="1F2937")
    font_bold = Font(name="Calibri", size=10, bold=True, color="1F2937")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.merge_cells("A1:G1")
    ws["A1"] = "BẢNG TỔNG HỢP THEO DÕI VÀ CHẤM CÔNG TĂNG CA (OT)"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:G2")
    display_period = f"Tháng {month_str}" if month_str != 'all' else "Toàn bộ thời gian"
    ws["A2"] = f"Hệ thống Phòng máy tính Tòa B4.2 (Cơ sở 2) • Ký theo dõi: {display_period}"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = align_center

    headers = ["STT", "Ngày OT", "Nhân sự", "Bộ phận / Vị trí", "Khung giờ & Số giờ", "Nội dung / Nhiệm vụ tăng ca", "Trạng thái"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[4].height = 26

    total_hours = 0.0
    for idx, r in enumerate(ot_rows, 1):
        row_num = 4 + idx
        h_val = float(r['hours'] or 0)
        total_hours += h_val
        st_text = "Đã hoàn thành (Đã làm)" if r['status'] == 'Completed' else "Chưa tới hạn"
        time_display = f"{r['start_time']} - {r['end_time']} ({h_val:g}h)" if (r['start_time'] and r['end_time']) else f"{h_val:g} giờ"

        cells_data = [
            (idx, align_center, font_body),
            (r['ot_date'], align_center, font_bold),
            (r['staff_name'], align_left, font_bold),
            (r['role'], align_left, font_body),
            (time_display, align_center, font_body),
            (r['purpose'] or "", align_left, font_body),
            (st_text, align_center, font_body)
        ]
        for c_idx, (val, al, fn) in enumerate(cells_data, 1):
            c = ws.cell(row=row_num, column=c_idx, value=val)
            c.alignment = al
            c.font = fn
            c.border = thin_border
        ws.row_dimensions[row_num].height = 20

    sum_row = 5 + len(ot_rows)
    ws.merge_cells(f"A{sum_row}:D{sum_row}")
    ws[f"A{sum_row}"] = "TỔNG CỘNG SỐ GIỜ TĂNG CA:"
    ws[f"A{sum_row}"].font = font_summary
    ws[f"A{sum_row}"].alignment = align_right
    ws[f"A{sum_row}"].fill = fill_summary

    sum_cell = ws.cell(row=sum_row, column=5, value=f"{total_hours:g} giờ")
    sum_cell.font = font_summary
    sum_cell.alignment = align_center
    sum_cell.fill = fill_summary

    ws.merge_cells(f"F{sum_row}:G{sum_row}")
    ws[f"F{sum_row}"].fill = fill_summary

    for c in range(1, 8):
        ws.cell(row=sum_row, column=c).border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['F'].width = 36

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Bang_Cham_Cong_OT_{month_str}.xlsx"
    return output, filename


@app.route('/admin/overtime/add', methods=['POST'])
@admin_required
def add_overtime():
    """Thêm một ca tăng ca (OT) mới."""
    ot_date = request.form.get('ot_date', '').strip()
    staff_name = request.form.get('staff_name', '').strip()
    role = request.form.get('role', 'Nhân viên Trực phòng').strip()
    start_time = request.form.get('start_time', '').strip()
    end_time = request.form.get('end_time', '').strip()
    raw_hours = request.form.get('hours', '').strip()
    purpose = request.form.get('purpose', '').strip()
    status = request.form.get('status', 'Pending').strip()
    notes = request.form.get('notes', '').strip()

    hours = 0.0
    if raw_hours:
        try:
            hours = float(raw_hours)
        except ValueError:
            hours = 0.0
    elif start_time and end_time:
        try:
            t1 = datetime.datetime.strptime(start_time, '%H:%M')
            t2 = datetime.datetime.strptime(end_time, '%H:%M')
            diff = (t2 - t1).total_seconds() / 3600.0
            if diff < 0:
                diff += 24.0
            hours = round(diff, 2)
        except Exception:
            hours = 0.0

    if not ot_date or not staff_name or hours <= 0:
        flash("⚠️ Vui lòng nhập đầy đủ Ngày, Tên nhân sự và Số giờ tăng ca hợp lệ!", "warning")
        return redirect(url_for('admin_panel', tab='overtime'))

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO overtime_records (ot_date, staff_name, role, start_time, end_time, hours, purpose, status, notes)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (ot_date, staff_name, role, start_time, end_time, hours, purpose, status, notes))
    conn.commit()
    conn.close()

    target_month = ot_date[:7] if len(ot_date) >= 7 else 'all'
    flash(f"✅ Đã ghi nhận lịch tăng ca cho {staff_name} ({hours:g} giờ - {ot_date}) thành công!", "success")
    return redirect(url_for('admin_panel', tab='overtime', ot_month=target_month))


@app.route('/admin/overtime/toggle-status/<int:ot_id>', methods=['POST'])
@admin_required
def toggle_overtime_status(ot_id):
    """Chuyển đổi trạng thái Đã làm <-> Chưa tới hạn cho ca tăng ca."""
    current_month = request.args.get('ot_month', '').strip()
    conn = get_db_connection()
    row = conn.execute("SELECT status, staff_name, ot_date FROM overtime_records WHERE id = ?", (ot_id,)).fetchone()
    if row:
        new_status = 'Completed' if row['status'] != 'Completed' else 'Pending'
        conn.execute("UPDATE overtime_records SET status = ? WHERE id = ?", (new_status, ot_id))
        conn.commit()
        st_text = "Đã làm" if new_status == 'Completed' else "Chưa tới hạn"
        flash(f"✅ Đã cập nhật trạng thái ca tăng ca của {row['staff_name']} sang '{st_text}'!", "success")
    conn.close()
    return redirect(url_for('admin_panel', tab='overtime', ot_month=current_month if current_month else None))


@app.route('/admin/overtime/delete/<int:ot_id>', methods=['POST'])
@admin_required
def delete_overtime(ot_id):
    """Xóa một bản ghi tăng ca."""
    current_month = request.args.get('ot_month', '').strip()
    conn = get_db_connection()
    conn.execute("DELETE FROM overtime_records WHERE id = ?", (ot_id,))
    conn.commit()
    conn.close()
    flash("🗑️ Đã xóa bản ghi tăng ca thành công!", "success")
    return redirect(url_for('admin_panel', tab='overtime', ot_month=current_month if current_month else None))


@app.route('/admin/overtime/export', methods=['GET'])
@admin_required
def export_overtime_excel():
    """Xuất file Excel bảng chấm công OT theo tháng."""
    month = request.args.get('month', datetime.date.today().strftime('%Y-%m')).strip()
    conn = get_db_connection()
    if month == 'all':
        rows = conn.execute("SELECT * FROM overtime_records ORDER BY ot_date ASC, id ASC").fetchall()
    else:
        rows = conn.execute("SELECT * FROM overtime_records WHERE ot_date LIKE ? ORDER BY ot_date ASC, id ASC", (f"{month}%",)).fetchall()
    conn.close()

    excel_buffer, filename = build_overtime_report_excel(rows, month)
    return send_file(
        excel_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route('/admin/approve-cancellation/<int:cancellation_id>', methods=['POST'])
@admin_required
def approve_cancellation(cancellation_id):
    """Cập nhật trạng thái thông báo nghỉ gấp (Đã tiếp nhận / Đã duyệt)."""
    new_status = request.form.get('status', 'Approved')
    if new_status not in ('Pending', 'Processing', 'Approved'):
        new_status = 'Approved'
    conn = get_db_connection()
    conn.execute("UPDATE cancellations SET status = ? WHERE id = ?", (new_status, cancellation_id))
    conn.commit()
    conn.close()
    
    status_label = "Đã tiếp nhận" if new_status == 'Processing' else ("Đã duyệt" if new_status == 'Approved' else "Chờ xử lý")
    flash(f"Đã cập nhật trạng thái báo nghỉ: {status_label}.", "success")
    return redirect(url_for('admin_panel', tab='cancellations'))

@app.route('/admin/delete-cancellation/<int:cancellation_id>', methods=['POST'])
@admin_required
def delete_cancellation(cancellation_id):
    """Xóa một thông báo nghỉ gấp."""
    conn = get_db_connection()
    conn.execute("DELETE FROM cancellations WHERE id = ?", (cancellation_id,))
    conn.commit()
    conn.close()
    flash("Đã xóa thông báo nghỉ thành công.", "success")
    return redirect(url_for('admin_panel', tab='cancellations'))

@app.route('/admin/bulk-delete-cancellations', methods=['POST'])
@admin_required
def bulk_delete_cancellations():
    """Xóa hàng loạt thông báo nghỉ gấp đã chọn."""
    selected_ids = [int(i) for i in request.form.getlist('selected_ids') if str(i).isdigit()]
    if selected_ids:
        conn = get_db_connection()
        placeholders = ','.join(['?'] * len(selected_ids))
        conn.execute(f"DELETE FROM cancellations WHERE id IN ({placeholders})", selected_ids)
        conn.commit()
        conn.close()
        flash(f"Đã xóa {len(selected_ids)} thông báo nghỉ được chọn.", "success")
    else:
        flash("Chưa có mục nào được chọn để xóa.", "warning")
    return redirect(url_for('admin_panel', tab='cancellations'))

@app.route('/admin/process-request/<int:request_id>', methods=['POST'])
@admin_required
def process_request(request_id):
    """Đánh dấu Đã tiếp nhận yêu cầu mượn phòng."""
    conn = get_db_connection()
    conn.execute("UPDATE booking_requests SET status = 'Processing' WHERE id = ?", (request_id,))
    conn.commit()
    conn.close()
    flash("Đã cập nhật trạng thái yêu cầu mượn phòng: Đã tiếp nhận.", "info")
    return redirect(url_for('admin_panel', tab='requests'))

@app.route('/admin/approve/<int:request_id>', methods=['POST'])
@admin_required
def approve_request(request_id):
    """Xử lý Duyệt yêu cầu mượn phòng (Lưu vết và thêm vào TKB nếu chưa thêm)."""
    conn = get_db_connection()
    req = conn.execute("SELECT * FROM booking_requests WHERE id = ?", (request_id,)).fetchone()
    
    if req:
        if req['status'] != 'Approved':
            class_code = req['class_code'] if req['class_code'] else 'MƯỢN PHÒNG'
            subject_name = req['subject_name'] if req['subject_name'] else req['purpose']
            department = req['department'] if req['department'] else req['requester_name']
            study_period = normalize_study_period_str(req['study_period']) if req['study_period'] else datetime.date.today().strftime('%d/%m/%Y')

            conn.execute('''
                INSERT INTO schedules (room, day_of_week, start_period, end_period, class_code, subject_name, department, study_period, num_weeks, is_ended_early)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, 0)
            ''', (req['room'], req['day_of_week'], req['start_period'], req['end_period'], class_code, subject_name, department, study_period))
        
        conn.execute("UPDATE booking_requests SET status = 'Approved' WHERE id = ?", (request_id,))
        conn.commit()
        flash(f"Đã duyệt yêu cầu của {req['requester_name']} vào phòng {req['room']}.", "success")
        
    conn.close()
    return redirect(url_for('admin_panel', tab='requests'))

@app.route('/admin/reject/<int:request_id>', methods=['POST'])
@admin_required
def reject_request(request_id):
    """Xử lý Từ chối yêu cầu mượn phòng."""
    admin_note = request.form.get('admin_note', '')
    conn = get_db_connection()
    conn.execute("UPDATE booking_requests SET status = 'Rejected', admin_note = ? WHERE id = ?", (admin_note, request_id))
    conn.commit()
    conn.close()
    
    flash("Đã từ chối yêu cầu mượn phòng.", "warning")
    return redirect(url_for('admin_panel', tab='requests'))

@app.route('/admin/delete-request/<int:request_id>', methods=['POST'])
@admin_required
def delete_request(request_id):
    """Xóa một yêu cầu mượn phòng khỏi danh sách."""
    conn = get_db_connection()
    conn.execute("DELETE FROM booking_requests WHERE id = ?", (request_id,))
    conn.commit()
    conn.close()
    flash("Đã xóa yêu cầu mượn phòng thành công.", "success")
    return redirect(url_for('admin_panel', tab='requests'))

@app.route('/admin/bulk-delete-requests', methods=['POST'])
@admin_required
def bulk_delete_requests():
    """Xóa hàng loạt yêu cầu mượn phòng đã chọn."""
    selected_ids = [int(i) for i in request.form.getlist('selected_ids') if str(i).isdigit()]
    if selected_ids:
        conn = get_db_connection()
        placeholders = ','.join(['?'] * len(selected_ids))
        conn.execute(f"DELETE FROM booking_requests WHERE id IN ({placeholders})", selected_ids)
        conn.commit()
        conn.close()
        flash(f"Đã xóa {len(selected_ids)} yêu cầu mượn phòng được chọn.", "success")
    else:
        flash("Chưa có mục nào được chọn để xóa.", "warning")
    return redirect(url_for('admin_panel', tab='requests'))

@app.route('/admin/schedule/<int:schedule_id>/toggle_early_end', methods=['POST'])
@admin_required
def toggle_early_end(schedule_id):
    """Đánh dấu Kết thúc môn sớm (Khóa cửa phòng & Giải phóng phòng trống) hoặc Khôi phục lại lịch."""
    return_tab = request.form.get('return_tab', 'schedules')
    room = request.form.get('room', '5.1')
    timetable_date = request.form.get('timetable_date', '')
    timetable_day = request.form.get('timetable_day', 'all')
    schedule_filter = request.form.get('schedule_filter', 'all')

    conn = get_db_connection()
    row = conn.execute("SELECT * FROM schedules WHERE id = ?", (schedule_id,)).fetchone()
    if row:
        current_ended = int(row['is_ended_early'] or 0)
        if current_ended == 0:
            today_str = datetime.date.today().strftime('%d/%m/%Y')
            conn.execute(
                "UPDATE schedules SET is_ended_early = 1, ended_early_date = ? WHERE id = ?",
                (today_str, schedule_id)
            )
            conn.commit()
            flash(f"🔒 Đã đánh dấu kết thúc sớm môn '{row['subject_name']}' (P.{row['room']}). Ca máy này đã chuyển sang Khóa cửa và giải phóng phòng trống!", "warning")
        else:
            conn.execute(
                "UPDATE schedules SET is_ended_early = 0, ended_early_date = NULL WHERE id = ?",
                (schedule_id,)
            )
            conn.commit()
            flash(f"🔓 Đã khôi phục lại lịch học môn '{row['subject_name']}' (P.{row['room']}) về trạng thái Mở cửa bình thường.", "success")
    conn.close()

    if return_tab == 'timetable':
        return redirect(url_for('admin_panel', tab='timetable', room=room, timetable_date=timetable_date, timetable_day=timetable_day))
    return redirect(url_for('admin_panel', tab='schedules', schedule_filter=schedule_filter))

@app.route('/admin/schedule/<int:schedule_id>/update_weeks', methods=['POST'])
@admin_required
def update_schedule_weeks(schedule_id):
    """Cập nhật số tuần lặp lại (num_weeks) cho một môn học cụ thể."""
    return_tab = request.form.get('return_tab', 'schedules')
    room = request.form.get('room', '5.1')
    timetable_date = request.form.get('timetable_date', '')
    timetable_day = request.form.get('timetable_day', 'all')
    schedule_filter = request.form.get('schedule_filter', 'all')
    new_weeks = request.form.get('num_weeks', 15, type=int)
    new_weeks = max(1, min(52, new_weeks))

    conn = get_db_connection()
    row = conn.execute("SELECT subject_name, room FROM schedules WHERE id = ?", (schedule_id,)).fetchone()
    if row:
        conn.execute("UPDATE schedules SET num_weeks = ? WHERE id = ?", (new_weeks, schedule_id))
        conn.commit()
        flash(f"Đã cập nhật chu kỳ lặp của môn '{row['subject_name']}' (P.{row['room']}) thành {new_weeks} tuần.", "success")
    conn.close()

    if return_tab == 'timetable':
        return redirect(url_for('admin_panel', tab='timetable', room=room, timetable_date=timetable_date, timetable_day=timetable_day))
    return redirect(url_for('admin_panel', tab='schedules', schedule_filter=schedule_filter))

@app.route('/admin/delete_schedule/<int:schedule_id>', methods=['POST'])
@admin_required
def delete_schedule(schedule_id):
    """Xóa một lịch học khỏi cơ sở dữ liệu."""
    conn = get_db_connection()
    conn.execute("DELETE FROM schedules WHERE id = ?", (schedule_id,))
    conn.commit()
    conn.close()
    
    flash("Đã xóa lịch học thành công.", "success")
    return redirect(url_for('admin_panel', tab='schedules'))

@app.route('/admin/delete_all_schedules', methods=['POST'])
@admin_required
def delete_all_schedules():
    """Xóa tất cả lịch học khỏi cơ sở dữ liệu."""
    conn = get_db_connection()
    conn.execute("DELETE FROM schedules")
    conn.commit()
    conn.close()
    
    flash("Đã xóa toàn bộ lịch học thành công.", "success")
    return redirect(url_for('admin_panel', tab='schedules'))

@app.route('/admin/add_schedule', methods=['POST'])
@admin_required
def add_schedule():
    """Xử lý form Thêm lịch mới thủ công của Admin."""
    room = request.form.get('room')
    day_of_week = int(request.form.get('day_of_week'))
    start_period = float(request.form.get('start_period'))
    end_period = float(request.form.get('end_period'))
    subject_name = request.form.get('subject_name')
    class_code = request.form.get('class_code')
    department = request.form.get('department')
    raw_study_period = request.form.get('study_period', '')
    study_period = normalize_study_period_str(raw_study_period)
    num_weeks = request.form.get('num_weeks', 15, type=int)
    num_weeks = max(1, min(52, num_weeks))
    
    if start_period >= end_period:
        flash("Tiết bắt đầu phải nhỏ hơn tiết kết thúc!", "danger")
        return redirect(url_for('admin_panel', tab='schedules'))
    if start_period < 1 or end_period > 10:
        flash("Tiết phải nằm trong khoảng từ 1 đến 10!", "danger")
        return redirect(url_for('admin_panel', tab='schedules'))

    conn = get_db_connection()
    conn.execute('''
        INSERT INTO schedules (room, day_of_week, start_period, end_period, class_code, subject_name, department, study_period, num_weeks, is_ended_early)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
    ''', (room, day_of_week, start_period, end_period, class_code, subject_name, department, study_period, num_weeks))
    
    conn.commit()
    conn.close()
    
    flash(f"Đã thêm lịch môn {subject_name} vào phòng {room} (lặp {num_weeks} tuần) thành công.", "success")
    return redirect(url_for('admin_panel', tab='schedules'))

import pandas as pd
from werkzeug.utils import secure_filename

@app.route('/admin/import_schedule', methods=['POST'])
@admin_required
def import_schedule():
    """Xử lý upload file Excel/CSV và import hàng loạt vào Database."""
    if 'file' not in request.files:
        flash("Không tìm thấy file tải lên.", "danger")
        return redirect(url_for('admin_panel', tab='schedules'))
        
    file = request.files['file']
    default_num_weeks = request.form.get('default_num_weeks', 15, type=int)
    default_num_weeks = max(1, min(52, default_num_weeks))
    
    if file.filename == '':
        flash("Chưa chọn file nào.", "warning")
        return redirect(url_for('admin_panel', tab='schedules'))

    if file:
        try:
            filename = secure_filename(file.filename)
            if filename.endswith('.csv'):
                df = pd.read_csv(file)
            elif filename.endswith(('.xls', '.xlsx')):
                df = pd.read_excel(file)
            else:
                flash("Định dạng file không được hỗ trợ. Chỉ nhận CSV hoặc Excel.", "danger")
                return redirect(url_for('admin_panel', tab='schedules'))

            df = df.fillna('')
            
            conn = get_db_connection()
            count = 0
            
            for index, row in df.iterrows():
                room = str(row.get('Phong', '')).strip()
                day_of_week = row.get('Thu', '')
                start_period = row.get('Tiet_BD', '')
                end_period = row.get('Tiet_KT', '')
                class_code = str(row.get('Ma_Lop', '')).strip()
                subject_name = str(row.get('Ten_Mon', '')).strip()
                department = str(row.get('Khoa', '')).strip()
                raw_period = row.get('Giai_Doan', '')
                study_period = normalize_study_period_str(raw_period)
                
                if not room or not day_of_week or start_period == '' or end_period == '':
                    continue

                try:
                    day_val = int(day_of_week)
                    start_val = float(start_period)
                    end_val = float(end_period)
                except ValueError:
                    continue

                # Kiểm tra trùng lịch (cùng phòng, thứ, tiết và cùng giai đoạn bắt đầu)
                existing = conn.execute('''
                    SELECT id FROM schedules
                    WHERE room = ? AND day_of_week = ? AND start_period = ? AND end_period = ? AND COALESCE(study_period, '') = ?
                ''', (room, day_val, start_val, end_val, study_period)).fetchone()
                if existing:
                    continue

                conn.execute('''
                    INSERT INTO schedules (room, day_of_week, start_period, end_period, class_code, subject_name, department, study_period, num_weeks, is_ended_early)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
                ''', (room, day_val, start_val, end_val, class_code, subject_name, department, study_period, default_num_weeks))
                count += 1
                
            conn.commit()
            conn.close()
            
            flash(f"Đã import thành công {count} dòng lịch học (chu kỳ lặp {default_num_weeks} tuần) vào hệ thống!", "success")
            
        except Exception as e:
            flash(f"Có lỗi xảy ra khi đọc file: {str(e)}", "danger")

    return redirect(url_for('admin_panel', tab='schedules'))

@app.route('/update_issue_status/<int:issue_id>', methods=['POST'])
@admin_required
def update_issue_status(issue_id):
    new_status = request.form.get('status')
    if new_status in ['Pending', 'Processing', 'Resolved', 'Approved', 'Đã tiếp nhận', 'Đã khắc phục', 'Đã duyệt']:
        if new_status == 'Đã tiếp nhận':
            new_status = 'Processing'
        elif new_status in ('Đã khắc phục', 'Đã duyệt', 'Approved'):
            new_status = 'Resolved'
        conn = get_db_connection()
        conn.execute('UPDATE issue_reports SET status = ? WHERE id = ?', (new_status, issue_id))
        conn.commit()
        conn.close()
        status_label = "Đã tiếp nhận" if new_status == 'Processing' else ("Đã duyệt" if new_status == 'Resolved' else "Chờ xử lý")
        flash(f"Đã cập nhật trạng thái sự cố: {status_label}.", "success")
    
    return redirect(url_for('admin_panel', tab='issues'))

@app.route('/admin/delete-issue/<int:issue_id>', methods=['POST'])
@admin_required
def delete_issue(issue_id):
    """Xóa một báo cáo sự cố khỏi danh sách."""
    conn = get_db_connection()
    conn.execute("DELETE FROM issue_reports WHERE id = ?", (issue_id,))
    conn.commit()
    conn.close()
    flash("Đã xóa báo cáo sự cố thành công.", "success")
    return redirect(url_for('admin_panel', tab='issues'))

@app.route('/admin/bulk-delete-issues', methods=['POST'])
@admin_required
def bulk_delete_issues():
    """Xóa hàng loạt báo cáo sự cố đã chọn."""
    selected_ids = [int(i) for i in request.form.getlist('selected_ids') if str(i).isdigit()]
    if selected_ids:
        conn = get_db_connection()
        placeholders = ','.join(['?'] * len(selected_ids))
        conn.execute(f"DELETE FROM issue_reports WHERE id IN ({placeholders})", selected_ids)
        conn.commit()
        conn.close()
        flash(f"Đã xóa {len(selected_ids)} báo cáo sự cố được chọn.", "success")
    else:
        flash("Chưa có mục nào được chọn để xóa.", "warning")
    return redirect(url_for('admin_panel', tab='issues'))


@app.route('/my-dashboard')
@login_required
def my_dashboard():
    """Trang cá nhân của User để xem lịch sử yêu cầu."""
    user_email = get_current_user_email()
        
    conn = get_db_connection()
    # Lấy booking requests
    my_bookings = conn.execute("SELECT * FROM booking_requests WHERE email = ? ORDER BY created_at DESC", (user_email,)).fetchall()
    # Lấy issue reports
    my_issues = conn.execute("SELECT * FROM issue_reports WHERE email = ? ORDER BY created_at DESC", (user_email,)).fetchall()
    # Lấy báo nghỉ
    my_cancellations = conn.execute("SELECT * FROM cancellations WHERE email = ? ORDER BY created_at DESC", (user_email,)).fetchall()
    conn.close()
    
    return render_template('my_dashboard.html', bookings=my_bookings, issues=my_issues, cancellations=my_cancellations)

if __name__ == '__main__':
    import sys
    app.config['TEMPLATES_AUTO_RELOAD'] = True
    is_pythonw = sys.executable.lower().endswith('pythonw.exe') or sys.stdout is None
    if is_pythonw:
        sys.stdout = open(os.devnull, 'w', encoding='utf-8')
        sys.stderr = open(os.devnull, 'w', encoding='utf-8')
    debug_mode = os.environ.get('FLASK_DEBUG', 'False').lower() == 'true'
    app.run(debug=debug_mode, use_reloader=(debug_mode and not is_pythonw), host='127.0.0.1', port=5000)
